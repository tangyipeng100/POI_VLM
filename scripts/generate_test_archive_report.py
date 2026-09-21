"""Generate a Chinese archival test report from the POI-VLM paper and raw logs."""

from __future__ import annotations

import csv
import hashlib
import json
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
PAPER = ROOT / "paper" / "poi_vlm_ei_paper.md"
FIGURES = ROOT / "paper" / "figures"
DATA150 = ROOT / "vlm_dataset_night_150"
DATA50 = ROOT / "vlm_dataset_night"
OUTPUT = ROOT / "output" / "doc" / "poi_vlm_ei_test_archive_report_cn.docx"

BLUE = "1F4E79"
LIGHT_BLUE = "EAF3F8"
PALE_BLUE = "F5FAFC"
LIGHT_GRAY = "F2F4F7"
GREEN = "EAF6EE"
ORANGE = "FFF2DC"
RED = "FCEBEC"
GRAY_TEXT = RGBColor(89, 89, 89)


def read_jsonl(path: Path) -> list[dict]:
    with path.open(encoding="utf-8") as handle:
        return [json.loads(line) for line in handle if line.strip()]


def read_csv(path: Path) -> list[dict]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest().upper()


def pct(num: int, den: int) -> str:
    return f"{num}/{den} ({num / den * 100:.1f}%)" if den else "-"


def fmt_num(value: float, digits: int = 2) -> str:
    return f"{value:.{digits}f}"


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 70, start: int = 90, bottom: int = 70, end: int = 90) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_text(cell, text: str, *, bold: bool = False, color: str | None = None, size: float = 9.2) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def set_table_borders(table, color: str = "B7C9D6", size: str = "6") -> None:
    tbl = table._tbl
    properties = tbl.tblPr
    borders = properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.9)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.15)
    section.right_margin = Cm(2.15)

    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.15

    for style_name, size, color in (("Title", 23, BLUE), ("Heading 1", 16, BLUE), ("Heading 2", 13, BLUE), ("Heading 3", 11, "365F91")):
        style = document.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(9 if style_name != "Title" else 0)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True

    if "Archive Note" not in document.styles:
        style = document.styles.add_style("Archive Note", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = document.styles["Normal"]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(9.5)
        style.font.color.rgb = GRAY_TEXT
        style.paragraph_format.left_indent = Cm(0.35)
        style.paragraph_format.right_indent = Cm(0.35)
        style.paragraph_format.space_after = Pt(4)

    if "Formula" not in document.styles:
        style = document.styles.add_style("Formula", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = document.styles["Normal"]
        style.font.name = "Cambria Math"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Cambria Math")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Cambria Math")
        style.font.size = Pt(10.5)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style.paragraph_format.space_after = Pt(4)

    add_header_footer(document.sections[0])


def add_header_footer(section) -> None:
    header = section.header
    header.is_linked_to_previous = False
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("POI-Guided VLM 户外机器人局部导航 | 测试存档报告")
    run.font.size = Pt(8)
    run.font.color.rgb = GRAY_TEXT
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("内部测试存档 · 第 ")
    run.font.size = Pt(8)
    run.font.color.rgb = GRAY_TEXT
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.append(field_begin)
    run._r.append(instruction)
    run._r.append(field_end)
    run2 = paragraph.add_run(" 页")
    run2.font.size = Pt(8)
    run2.font.color.rgb = GRAY_TEXT


def add_paragraph(document, text: str = "", *, bold_prefix: str | None = None, style: str | None = None, align=None):
    paragraph = document.add_paragraph(style=style)
    if align is not None:
        paragraph.alignment = align
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        first.bold = True
        paragraph.add_run(text[len(bold_prefix):])
    else:
        paragraph.add_run(text)
    return paragraph


def add_bullet(document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.add_run(text)


def add_numbered(document, number: int, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.4)
    paragraph.paragraph_format.first_line_indent = Cm(-0.4)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.add_run(f"{number}. ").bold = True
    paragraph.add_run(text)


def add_callout(document, title: str, body: str, fill: str = LIGHT_BLUE) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color="B7C9D6", size="6")
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(title)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(BLUE)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    paragraph.add_run("\n" + body)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(document, headers: list[str], rows: list[list[str]], *, font_size: float = 8.9, header_fill: str = BLUE):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, value in enumerate(headers):
        set_cell_text(header.cells[index], value, bold=True, color="FFFFFF", size=font_size)
        set_cell_shading(header.cells[index], header_fill)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        fill = PALE_BLUE if row_index % 2 == 0 else "FFFFFF"
        for index, value in enumerate(row):
            set_cell_text(cells[index], value, size=font_size)
            set_cell_shading(cells[index], fill)
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_image(document, path: Path, caption: str, *, width: float = 6.35) -> None:
    if not path.exists():
        add_paragraph(document, f"[图像缺失：{path}]", style="Archive Note")
        return
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    caption_paragraph = document.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.space_after = Pt(6)
    run = caption_paragraph.add_run(caption)
    run.italic = True
    run.font.size = Pt(8.5)
    run.font.color.rgb = GRAY_TEXT


def add_formula(document, text: str) -> None:
    paragraph = document.add_paragraph(style="Formula")
    run = paragraph.add_run(text)
    run.font.name = "Cambria Math"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Cambria Math")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Cambria Math")


def add_page_break(document: Document) -> None:
    document.add_page_break()


def metric_counts(rows: list[dict]) -> dict[str, int]:
    return {
        "n": len(rows),
        "action": sum(bool(row.get("action_correct")) for row in rows),
        "target": sum(bool(row.get("target_correct")) for row in rows),
        "relaxed_action": sum(bool(row.get("relaxed_action_correct")) for row in rows),
        "relaxed_target": sum(bool(row.get("relaxed_target_correct")) for row in rows),
    }


def quantile(values: list[float], ratio: float) -> float:
    ordered = sorted(values)
    index = (len(ordered) - 1) * ratio
    low = int(index)
    high = min(low + 1, len(ordered) - 1)
    fraction = index - low
    return ordered[low] * (1 - fraction) + ordered[high] * fraction


def target_display(row: dict, prefix: str = "pred") -> str:
    action = row.get(f"{prefix}_action") or row.get("action")
    if action == "go_to_poi":
        value = row.get(f"{prefix}_poi_number") or row.get("poi_number")
        return str(value) if value not in (None, "") else "-"
    if action == "go_between_pois":
        values = row.get(f"{prefix}_poi_numbers")
        if values is None:
            values = row.get("poi_numbers")
        if isinstance(values, str):
            try:
                values = json.loads(values)
            except json.JSONDecodeError:
                pass
        if values:
            return "[" + ", ".join(str(value) for value in values) + "]"
        return "-"
    if action == "rotate":
        direction = row.get(f"{prefix}_rotate_direction") or row.get("rotate_direction") or "?"
        angle = row.get(f"{prefix}_rotate_angle_deg") or row.get("rotate_angle_deg")
        return f"{direction}/{angle}°" if angle not in (None, "") else str(direction)
    return "-"


def label_target_display(row: dict) -> str:
    action = row.get("label_action") or row.get("action")
    if action == "go_to_poi":
        value = row.get("label_poi_number") or row.get("poi_number")
        return str(value) if value not in (None, "") else "-"
    if action == "rotate":
        direction = row.get("label_rotate_direction") or row.get("rotate_direction") or "?"
        angle = row.get("label_rotate_angle_deg") or row.get("rotate_angle_deg")
        return f"{direction}/{angle}°" if angle not in (None, "") else str(direction)
    return "-"


def source_row(relative: str, role: str) -> list[str]:
    path = ROOT / relative
    return [relative, role, f"{path.stat().st_size:,}", sha256(path)[:16] + "…"]


def set_landscape(section) -> None:
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Cm(1.4)
    section.bottom_margin = Cm(1.4)
    section.left_margin = Cm(1.25)
    section.right_margin = Cm(1.25)


def main() -> None:
    annotations = read_csv(DATA150 / "annotations.csv")
    gpt_rows = read_jsonl(DATA150 / "action_predictions.jsonl")
    qwen_rows = read_jsonl(DATA150 / "qwen_h800_eval" / "action_predictions.jsonl")
    samples150 = read_jsonl(DATA150 / "samples.jsonl")
    samples50 = read_jsonl(DATA50 / "samples.jsonl")
    comparison = json.loads((DATA150 / "qwen_h800_eval" / "qwen_vs_gpt55_comparison.json").read_text(encoding="utf-8"))

    annotation_counts = Counter(row.get("action") for row in annotations)
    gpt_metrics = metric_counts(gpt_rows)
    qwen_metrics = metric_counts(qwen_rows)
    qwen_latencies = [float(row["latency_s"]) for row in qwen_rows if row.get("latency_s") is not None]

    sample_poi_counts = [len(row.get("pois", [])) for row in samples150]
    route = [row.get("route", {}) for row in samples150]
    route_total = route[0].get("total_m", 0.0)
    passed_values = [item.get("passed_m", 0.0) for item in route]
    remaining_values = [item.get("remaining_m", 0.0) for item in route]
    heading_values = [item.get("heading_deg", 0.0) for item in route]
    distance_values = [item.get("distance_to_route_m", 0.0) for item in route]
    poi_edge_counts = Counter(point.get("edge_type") for sample in samples150 for point in sample.get("pois", []))

    qwen_by_id = {row["sample_id"]: row for row in qwen_rows}
    gpt_by_id = {row["sample_id"]: row for row in gpt_rows}

    document = Document()
    configure_document(document)

    # Cover
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(55)
    title.paragraph_format.space_after = Pt(10)
    run = title.add_run("POI-Guided VLM\n户外机器人局部导航测试存档报告")
    run.bold = True
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor.from_string(BLUE)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(20)
    run = subtitle.add_run("基于论文正文、标注文件与模型逐帧预测日志的中文归档版本")
    run.font.size = Pt(12)
    run.font.color.rgb = GRAY_TEXT

    date_line = document.add_paragraph()
    date_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_line.paragraph_format.space_after = Pt(7)
    date_run = date_line.add_run(f"生成日期：{datetime.now().strftime('%Y-%m-%d')}")
    date_run.font.size = Pt(9.5)
    date_run.font.color.rgb = GRAY_TEXT

    cover_table = add_table(
        document,
        ["项目", "归档内容"],
        [
            ["论文题目", "POI-Guided Vision-Language Decision Making for Outdoor Robot Local Navigation"],
            ["作者", "Yipeng Tang†, Qiancheng Ye†, Zhixing Song, Haoge Jiang, Chaochao He, Zhenyu Xu, Jinfei Gao, Baoping Cheng, Ting Huang, Qiran Pu*"],
            ["作者标记", "† 共一作者；* 通讯作者"],
            ["报告性质", "中文测试存档；用于记录论文方法、数据、结果、原始日志与可复核性边界"],
        ],
        font_size=9.0,
    )
    for row in cover_table.rows[1:]:
        set_cell_shading(row.cells[0], LIGHT_BLUE)
        row.cells[0].paragraphs[0].runs[0].bold = True

    add_page_break(document)
    add_callout(
        document,
        "归档口径",
        "本报告不新增实验结论。论文中的 50 帧日间结果按论文表格记录；150 帧夜间结果使用当前工作区中的 annotations.csv、GPT-5.5 与 Qwen-H800 逐帧 JSONL 重新统计。凡当前目录缺少原始标注或硬件记录的内容，均明确标注为“论文记录”或“未提供”，不作推断。",
        fill=ORANGE,
    )

    # Executive summary
    document.add_heading("执行摘要", level=1)
    add_paragraph(
        document,
        "本文研究一个面向户外移动机器人的离线局部导航决策框架。系统先利用 PP-LiteSeg 风格的可通行区域分割模型获得前视可行区域，再通过多边形边界逻辑生成少量、可解释的候选兴趣点（POI）；同时将采集到的 GPS 轨迹渲染为俯视路线图，最后把前视 POI 图和俯视路线图拼接后交给视觉语言模型（VLM）选择下一步动作。模型输出被约束为结构化 JSON，因此可以逐帧解析、计分和复核。"
    )
    add_paragraph(
        document,
        "论文的实验边界是离线帧级决策评估，不是闭环机器人导航实验。150 帧夜间测试中，GPT-5.5 的动作准确率为 92.0%，严格目标准确率为 64.7%，宽松目标准确率为 66.7%；Qwen-H800 的对应结果为 91.3%、78.7% 和 78.7%。这说明两个模型通常能够判断动作类型，但精确选中人工 POI 目标仍受候选点歧义、少数旋转/跳过样本和 GPS 视图偏差影响。"
    )
    add_paragraph(
        document,
        "从工程归档角度看，当前资料已经能复核 150 帧主测试和两个模型的逐帧结果；仍需注意：论文表中的 50 帧结果在当前目录没有对应的独立标注文件，GPT 模型硬件/接口成本没有记录，分割模型也没有独立 IoU 等指标。因此，本报告适合保存方法与测试证据，不应被解读为闭环导航性能或通用模型能力证明。"
    )

    add_table(
        document,
        ["项目", "存档结论"],
        [
            ["任务形态", "前视可行性 + GPS 路线趋势条件下的离线局部动作选择"],
            ["主测试集", "150 帧；论文称夜间集；当前可直接复核"],
            ["动作空间", "go_to_poi、go_between_pois、rotate、skip"],
            ["主结果", "GPT-5.5 动作 138/150；Qwen-H800 动作 137/150"],
            ["最明显差异", "Qwen-H800 严格目标 118/150，高于 GPT-5.5 的 97/150"],
            ["尚未证明", "闭环控制、真实机器人成功率、避障保证、泛化优势"],
        ],
        font_size=9.1,
    )

    # 1 scope
    document.add_heading("1. 报告范围与资料版本", level=1)
    add_paragraph(
        document,
        "本报告覆盖论文的研究问题、方法模块、数据集与标注方式、评价指标、日间/夜间结果、逐模型误差、定位偏差探针、参考文献验证状态和原始文件校验信息。报告以工作区中的论文 Markdown 为叙述基准，以 150 帧数据目录中的 CSV/JSONL 为主测试证据。"
    )
    add_table(
        document,
        ["资料", "作用", "当前归档状态"],
        [
            ["paper/poi_vlm_ei_paper.md", "论文正文、方法、表格、结论与参考文献", "存在；SHA-256 见附录"],
            ["vlm_dataset_night_150/annotations.csv", "150 帧人工动作与目标标注", "存在；150 行"],
            ["vlm_dataset_night_150/action_predictions.jsonl", "GPT-5.5 逐帧结构化预测及正确性字段", "存在；150 行"],
            ["vlm_dataset_night_150/qwen_h800_eval/action_predictions.jsonl", "Qwen-H800 逐帧预测及时延", "存在；150 行"],
            ["vlm_dataset_night_150/samples.jsonl", "150 帧样本、POI、GPS 路线与图像路径", "存在；150 行"],
            ["paper/reference_verification_report.md", "参考文献逐条核验记录", "存在；论文报告称无不可核验文献"],
        ],
        font_size=8.8,
    )
    add_callout(
        document,
        "命名与可复核性说明",
        "项目目录名 vlm_dataset_night 与 vlm_dataset_night_150 属于历史命名。论文把 50 帧子集称为日间集、150 帧子集称为夜间集；本报告沿用论文口径，不根据目录名重新判断采集时段。当前工作区能独立复算 150 帧结果，但没有发现 50 帧对应的 annotations.csv/annotations.json，因此 50 帧准确率只作为论文已报告结果保存。",
        fill=ORANGE,
    )

    # 2 paper summary
    document.add_heading("2. 论文内容总结", level=1)
    document.add_heading("2.1 研究问题与边界", level=2)
    add_paragraph(
        document,
        "户外机器人局部导航需要同时回答两个问题：前方哪些区域在视觉上可通行，以及下一步方向是否符合更高层路线意图。单独依赖前视图可能走向路线不一致的区域；单独依赖 GPS 线路又可能受到定位漂移、视图不同步或暂时遮挡影响。论文提出的中间层方案不让 VLM 直接输出连续控制量，而是让它在由可通行区域约束的 POI 候选中选择一个可解释动作。"
    )
    document.add_heading("2.2 方法流水线", level=2)
    add_image(document, FIGURES / "fig1_framework_editable_arrowfix.jpg", "图 1  论文方法总流程（原论文图，英文标签保留）", width=6.25)
    for number, text in enumerate(
        [
            "输入户外前视 RGB 视频帧和 GPS 轨迹记录。",
            "使用 PP-LiteSeg 风格二值分割模型估计可通行区域。论文记录训练图像约 5000 张，均有人工二值可通行标注。",
            "对分割掩膜做形态学清理、连通域筛选和最大轮廓提取，再进行自适应多边形近似。",
            "跳过通常对应车体近场边界的 bottom 边，在有效 left/right/top 边上取中点，并增加最小 y 值顶点作为 vanishing POI。",
            "对候选 POI 做优先级 NMS，并限制显示数量，使 VLM 面对少量可解释目标。",
            "将 GPS 当前点、前进方向、已通过轨迹、剩余轨迹、起点和终点绘制到俯视图。",
            "将前视 POI 图与俯视路线图组成双视角输入，调用 VLM，解析结构化 JSON 并与人工标签比较。",
        ],
        start=1,
    ):
        add_numbered(document, number, text)
    add_image(document, FIGURES / "fig4_question_image.jpg", "图 2  双视角 VLM 输入示例：左侧为前视 POI，右侧为 GPS 俯视路线。", width=6.25)

    document.add_heading("2.3 动作空间与输出字段", level=2)
    add_table(
        document,
        ["动作代码", "中文含义", "目标字段", "使用条件"],
        [
            ["go_to_poi", "前往一个 POI", "poi_number", "前视可行且与路线方向一致"],
            ["go_between_pois", "前往两个 POI 之间的中点", "poi_numbers + interpolation", "中间区域比单个 POI 更合适"],
            ["rotate", "先旋转再重新选择目标", "rotate_direction + rotate_angle_deg", "视角未对准路线或没有合适前向目标"],
            ["skip", "跳过当前样本", "无", "图像不可用或没有可用候选"],
        ],
        font_size=8.9,
    )
    add_paragraph(document, "论文中的核心数据关系可归纳为：", bold_prefix="论文中的核心数据关系可归纳为：")
    add_formula(document, "I_i -> M_i = S_phi(I_i) -> P_i = G(M_i) -> Q_i = (I_i^front, I_i^map, P_i, R_i) -> y_hat_i = f_theta(Q_i, q)")
    add_formula(document, "P_i = G(M_i) = {p_ij}_{j=1}^{K_i},    p_ij = (n_ij, u_ij, v_ij, c_ij, s_ij)")
    add_paragraph(document, "这里的 POI 是由分割结果约束出来的视觉导航锚点，而不是直接给控制器执行的连续坐标。")

    # 3 data and annotation
    document.add_heading("3. 测试数据与标注归档", level=1)
    document.add_heading("3.1 论文报告的数据组成", level=2)
    add_table(
        document,
        ["数据组件", "用途", "样本量", "人工标签/元数据"],
        [
            ["分割训练集", "训练 PP-LiteSeg 风格可通行区域分割器", "约 5000 张", "二值可通行区域掩膜；精确文件未随当前工作区归档"],
            ["日间 VLM 评估集", "离线动作决策评估", "50 帧", "论文表 1：44 go_to_poi、6 rotate、0 skip"],
            ["夜间 VLM 评估集", "离线动作决策评估", "150 帧", "143 go_to_poi、3 rotate、4 skip；当前可独立复核"],
        ],
        font_size=8.9,
    )
    add_image(document, FIGURES / "fig5_dataset_label_distribution.png", "图 3  论文报告的日间与夜间动作标签分布。", width=5.85)

    document.add_heading("3.2 150 帧主测试集的原始数据统计", level=2)
    add_table(
        document,
        ["字段", "提取结果", "说明"],
        [
            ["样本数", "150", "samples.jsonl、annotations.csv、两个模型 JSONL 均为 150 条"],
            ["图像尺寸", "1920 × 1080", "150 帧一致"],
            ["样本 frame_index", "0–842", "抽帧索引，不是连续视频帧"],
            ["GPS frame", "5–10110", "原始 GPS 记录索引范围"],
            ["路线来源", "gps_track（150/150）", "直接使用采集到的 GPS 轨迹渲染"],
            ["总路线长度", f"{fmt_num(route_total, 2)} m", "150 帧记录一致"],
            ["passed_m", f"{fmt_num(min(passed_values), 2)}–{fmt_num(max(passed_values), 2)} m", "当前样本在路线上的已通过距离"],
            ["remaining_m", f"{fmt_num(min(remaining_values), 2)}–{fmt_num(max(remaining_values), 2)} m", "当前样本的剩余路线距离"],
            ["distance_to_route_m", f"{fmt_num(min(distance_values), 2)}–{fmt_num(max(distance_values), 2)} m", "当前 150 帧均为 0.00 m"],
            ["heading_deg", f"{fmt_num(min(heading_values), 2)}–{fmt_num(max(heading_values), 2)}°", "lookahead 路线方向"],
            ["POI 数量/帧", f"{min(sample_poi_counts)}–{max(sample_poi_counts)}；平均 {fmt_num(sum(sample_poi_counts) / len(sample_poi_counts), 2)}", "由多边形边界生成"],
            ["POI 边类型计数", f"vanishing {poi_edge_counts['vanishing']}；left {poi_edge_counts['left']}；right {poi_edge_counts['right']}；top {poi_edge_counts['top']}", "跨 150 帧累计候选点计数"],
        ],
        font_size=8.5,
    )

    document.add_heading("3.3 标注规则与标注结果", level=2)
    add_paragraph(document, "150 帧人工标签来自 annotations.csv。当前统计为：")
    add_table(
        document,
        ["人工动作", "数量", "占比", "标签解释"],
        [
            ["go_to_poi", str(annotation_counts["go_to_poi"]), f"{annotation_counts['go_to_poi'] / len(annotations) * 100:.1f}%", "选择一个可行 POI"],
            ["rotate", str(annotation_counts["rotate"]), f"{annotation_counts['rotate'] / len(annotations) * 100:.1f}%", "先旋转；当前样本较少"],
            ["skip", str(annotation_counts["skip"]), f"{annotation_counts['skip'] / len(annotations) * 100:.1f}%", "跳过坏帧或不可评估帧"],
        ],
        font_size=9.0,
    )
    add_paragraph(
        document,
        "标注工具支持多 POI 可接受集合，但当前 150 帧 CSV 中大多数记录是单个 poi_number。由于 rotate 与 skip 仅分别有 3 帧和 4 帧，它们的类别准确率只能作为诊断信息，不能视为稳定的泛化估计。"
    )

    document.add_heading("3.4 测试输入与数据链", level=2)
    add_table(
        document,
        ["阶段", "输入", "输出", "归档文件/图像"],
        [
            ["分割", "前视 RGB 帧", "二值 drivable mask", "论文图 2；训练集原始掩膜未完整归档"],
            ["POI 生成", "drivable mask", "带编号候选 POI", "generated_front_poi/*.jpg；samples.jsonl 中记录坐标"],
            ["路线渲染", "GPS 轨迹 + 当前帧", "俯视 route map", "route_maps/*.jpg；route_source=gps_track"],
            ["VLM 输入", "前视 POI + route map + prompt", "双视角 question image", "question_images/*.jpg"],
            ["人工评估", "question image", "动作/目标标签", "annotations.csv/.json/.jsonl"],
            ["模型评估", "同一 question image 和 prompt", "结构化 JSON 预测", "action_predictions.jsonl；Qwen 对比目录"],
        ],
        font_size=8.6,
    )

    # 4 metrics
    document.add_heading("4. 评价指标", level=1)
    add_paragraph(document, "设第 i 个样本的人工动作为 a_i，人工可接受目标集合为 T_i，模型预测动作和目标为 a_hat_i、T_hat_i。论文采用以下三类指标：")
    add_formula(document, "动作准确率 = 动作类型预测正确的样本数 / N")
    add_formula(document, "严格目标准确率 = 动作类型正确且目标也正确的样本数 / N")
    add_formula(document, "宽松目标准确率 = 严格目标匹配数 + 可接受的两点中间目标匹配数 / N")
    add_paragraph(document, "严格目标指标要求模型既选对动作类型又选中人工 POI；宽松目标指标允许 go_between_pois 的两个端点中有一个属于人工可接受目标集合。skip 的目标字段为空，但动作类型正确即可计入目标正确。")
    add_callout(
        document,
        "分母注意事项",
        "主表中的 Target Acc. 以全部 N 帧为分母；按动作类型拆分时，POI 子集以 143 帧为分母，rotate 以 3 帧为分母，skip 以 4 帧为分母。因此“100/150”和“100/143”表达的是不同口径，不能直接互换。",
        fill=LIGHT_BLUE,
    )

    # 5 results
    document.add_heading("5. 测试结果", level=1)
    document.add_heading("5.1 主结果对比", level=2)
    add_table(
        document,
        ["设置", "样本", "预测分布", "动作准确率", "严格目标准确率", "宽松目标准确率", "证据来源"],
        [
            ["日间，GPT-style", "50", "47 go_to_poi；3 rotate", "45/50 (90.0%)", "34/50 (68.0%)", "34/50 (68.0%)", "论文表 2；当前无独立标签文件"],
            ["夜间，GPT-5.5", "150", "145 go_to_poi；3 go_between_pois；2 rotate", pct(gpt_metrics["action"], 150), pct(gpt_metrics["target"], 150), pct(gpt_metrics["relaxed_target"], 150), "action_predictions.jsonl；已复算"],
            ["夜间，Qwen-H800", "150", "142 go_to_poi；7 rotate；1 skip", pct(qwen_metrics["action"], 150), pct(qwen_metrics["target"], 150), pct(qwen_metrics["relaxed_target"], 150), "qwen_h800_eval JSONL；已复算"],
        ],
        font_size=7.8,
    )
    add_paragraph(
        document,
        f"补充记录：GPT-5.5 的原始 JSONL 还给出宽松动作匹配 {gpt_metrics['relaxed_action']}/150 ({gpt_metrics['relaxed_action'] / 150 * 100:.1f}%)；Qwen-H800 的宽松动作匹配为 {qwen_metrics['relaxed_action']}/150 ({qwen_metrics['relaxed_action'] / 150 * 100:.1f}%)。该字段用于保存原始评测日志，论文主表没有单独列出。"
    )
    add_image(document, FIGURES / "fig7_accuracy_comparison.png", "图 4  论文报告的离线决策准确率对比。", width=5.9)

    document.add_heading("5.2 按动作类型的结果", level=2)
    add_table(
        document,
        ["模型", "POI 严格", "POI 宽松", "rotate 目标", "skip 动作/目标", "主要含义"],
        [
            ["GPT-5.5", "97/143 (67.8%)", "100/143 (69.9%)", "0/3 (0.0%)", "0/4 (0.0%)", "POI 目标存在歧义；罕见 rotate/skip 易被判为前往 POI"],
            ["Qwen-H800", "117/143 (81.8%)", "117/143 (81.8%)", "1/3 (33.3%)", "0/4 (0.0%)", "POI 匹配更高；仍有过度旋转和 skip 识别问题"],
        ],
        font_size=8.2,
    )
    add_image(document, FIGURES / "fig8_action_type_breakdown.png", "图 5  夜间测试按动作类型的性能分解；柱顶数字为正确数/该类样本数。", width=5.9)

    document.add_heading("5.3 Qwen-H800 推理时延与模型对比", level=2)
    add_table(
        document,
        ["时延字段", "提取值", "说明"],
        [
            ["平均时延", f"{fmt_num(sum(qwen_latencies) / len(qwen_latencies), 4)} s", "150 条逐帧日志"],
            ["P50", f"{fmt_num(quantile(qwen_latencies, 0.50), 4)} s", "中位数"],
            ["P95", f"{fmt_num(quantile(qwen_latencies, 0.95), 4)} s", "95 分位"],
            ["最小/最大", f"{fmt_num(min(qwen_latencies), 4)} / {fmt_num(max(qwen_latencies), 4)} s", "逐帧范围"],
            ["150 帧累计", f"{fmt_num(sum(qwen_latencies), 3)} s", "仅为日志累计推理时延"],
        ],
        font_size=8.9,
    )
    paired = comparison.get("paired", {})
    add_table(
        document,
        ["配对结果", "数量", "含义"],
        [
            ["两者都正确", str(paired.get("both_correct", "-")), "同一帧 GPT 与 Qwen 均命中评价条件"],
            ["两者都错误", str(paired.get("both_wrong", "-")), "同一帧两者均未命中"],
            ["Qwen 正确、GPT 错误", str(paired.get("qwen_correct_gpt_wrong", "-")), "Qwen 的独占正确帧"],
            ["GPT 正确、Qwen 错误", str(paired.get("gpt_correct_qwen_wrong", "-")), "GPT 的独占正确帧"],
            ["预测存在差异", str(comparison.get("prediction_disagreements", "-")), "两个模型的结构化输出不同"],
        ],
        font_size=8.9,
    )
    add_callout(
        document,
        "硬件成本记录边界",
        "当前工作区有 Qwen-H800 标签和逐帧 latency_s，但没有 GPT-5.5 的调用时延、GPU 型号/显存、输入输出 token、API 价格或单帧电费记录；Qwen 也没有完整的硬件占用/费用日志。因此本报告只归档可见时延，不把“Qwen-H800”解释成完整计算成本。",
        fill=ORANGE,
    )

    document.add_heading("5.4 代表性案例与 GPS 偏差探针", level=2)
    add_image(document, FIGURES / "fig6_case_studies.jpg", "图 6  论文代表性案例：正确样本、POI 目标失败样本和定位偏差探针。", width=6.15)
    add_paragraph(document, "普通失败案例中，人工目标为 POI 5，模型选择 POI 4。两者在前视图中都可能看起来可行，因此该错误更接近细粒度目标歧义，而不是动作类型完全错误。")
    add_image(document, FIGURES / "fig7_localization_bias_case.jpg", "图 7  定位偏差探针：前视帧 frame_000384 配对了 route frame_000305 的俯视路线，模型仍选择 POI 1。", width=6.15)
    add_paragraph(document, "该探针是一个定性 hard case，不是系统性鲁棒性基准。它说明提示词中关于 GPS 可能存在噪声/偏差的说明，在这一帧上没有阻止模型利用较大的路线趋势和前视可行区域；但仅凭 1 个配对样本不能量化定位误差鲁棒性。")

    # 6 errors and limitations
    document.add_heading("6. 错误分析与研究限制", level=1)
    document.add_heading("6.1 动作混淆记录", level=2)
    gpt_conf = defaultdict(Counter)
    for row in gpt_rows:
        gpt_conf[row.get("label_action")][row.get("pred_action")] += 1
    qwen_conf = defaultdict(Counter)
    for row in qwen_rows:
        qwen_conf[row.get("label_action")][row.get("pred_action")] += 1
    conf_rows = []
    for label in ["go_to_poi", "rotate", "skip"]:
        conf_rows.append(
            [
                label,
                ", ".join(f"{key}:{value}" for key, value in sorted(gpt_conf[label].items())),
                ", ".join(f"{key}:{value}" for key, value in sorted(qwen_conf[label].items())),
            ]
        )
    add_table(document, ["人工标签", "GPT-5.5 预测分布", "Qwen-H800 预测分布"], conf_rows, font_size=8.8)
    add_paragraph(document, "关键观察：")
    for text in [
        "GPT-5.5 对 143 个 go_to_poi 标签中有 138 个保持动作类型正确，但其中 3 个输出 go_between_pois、2 个输出 rotate；3 个 rotate 和 4 个 skip 全部被判成 go_to_poi。",
        "Qwen-H800 对 143 个 go_to_poi 标签中有 136 个动作类型正确，6 个被判为 rotate，1 个被判为 skip；3 个 rotate 中命中 1 个，4 个 skip 全部被判为 go_to_poi。",
        "Qwen 的主要优势在 POI 目标匹配，而不是稀有动作的识别；稀有类别样本太少，不能据此做稳定的类别泛化结论。",
    ]:
        add_bullet(document, text)

    document.add_heading("6.2 论文明确的限制", level=2)
    for text in [
        "只进行了离线帧级决策评估，没有展示闭环机器人导航、真实控制成功率、避障保证或控制稳定性。",
        "路线数量、天气和光照变化有限；rotate 与 skip 类别严重不平衡。",
        "当前没有传统局部规划器 baseline，因此不能声称相对于传统方法有导航性能优势。",
        "没有报告 PP-LiteSeg 独立分割指标，分割误差与 VLM 决策误差尚未分离。",
        "单一人工目标可能低估多个合理 POI 的情况，后续应继续完善多标签目标集合。",
        "GPS 俯视图可能存在定位误差或与前视图不同步；当前只做了一个定性探针。",
    ]:
        add_bullet(document, text)

    document.add_heading("6.3 归档中应保留的待补齐项", level=2)
    add_table(
        document,
        ["待补齐项", "当前状态", "对复核的影响"],
        [
            ["50 帧子集原始人工标注", "当前目录未发现独立 annotations 文件", "无法从当前目录独立重算 45/50、34/50"],
            ["GPT-5.5 推理成本", "未记录硬件、token、接口账单或时延", "只能复核准确率，不能复核算力/费用"],
            ["PP-LiteSeg 分割指标", "论文只记录约 5000 张训练图像", "无法单独判断前端掩膜质量"],
            ["系统性 GPS 漂移实验", "只有一个定性定位偏差探针", "不能量化不同漂移量下的性能曲线"],
            ["真实闭环验证", "未开展", "论文结论只限离线决策接口"],
        ],
        font_size=8.6,
    )

    # 7 conclusion and reference status
    document.add_heading("7. 归档结论", level=1)
    for text in [
        "论文把连续局部导航问题转化为可视化、可解析、可逐帧计分的 POI 决策问题。",
        "150 帧主测试的数据链完整：样本、GPS 路线、POI、人工标签、GPT-5.5 预测和 Qwen-H800 预测均可在工作区找到。",
        "两个模型的动作类型准确率都超过 91%；Qwen-H800 在当前数据上的严格 POI 目标匹配高于 GPT-5.5。",
        "主要剩余问题是细粒度 POI 歧义、少量 rotate/skip 样本、GPS 偏差的系统性量化和闭环控制验证。",
        "本报告适合作为论文测试证据和版本存档；对外表述时应保留“offline decision evaluation”边界。",
    ]:
        add_bullet(document, text)

    document.add_heading("8. 参考文献核验存档", level=1)
    add_paragraph(
        document,
        "论文共列出 10 条参考文献。工作区的 reference_verification_report.md 记录了逐条核验来源。该报告记载：参考文献 [1]、[2]、[5]、[6]、[7]、[8]、[9]、[10] 的题名与书目信息已核对；[3] 修正了原有题名/作者不匹配；[4] 补全了作者并修正题名。按现有核验记录，没有剩余不可核验或虚构文献。"
    )
    add_table(
        document,
        ["核验范围", "数量", "状态"],
        [
            ["论文参考文献总数", "10", "已列入正文"],
            ["需要实质元数据修正", "2 条（[3]、[4]）", "已在论文版本中修正"],
            ["核验报告", "paper/reference_verification_report.md", "存在；记录公开来源链接"],
            ["预印本比例提醒", "按投稿方规则另行核对", "本报告不重新修改论文参考文献列表"],
        ],
        font_size=8.8,
    )

    # Appendices in portrait first.
    add_page_break(document)
    document.add_heading("附录 A  数据字段说明", level=1)
    add_table(
        document,
        ["字段", "所在文件", "含义"],
        [
            ["sample_id", "samples.jsonl / JSONL", "样本编号，如 frame_000384"],
            ["frame_index", "samples.jsonl", "抽帧序号"],
            ["gps / route", "samples.jsonl", "WGS84/GCJ-02 位置、路线投影、heading、passed、remaining"],
            ["pois", "samples.jsonl", "候选 POI 编号、像素位置、边类型和颜色"],
            ["action / poi_number", "annotations.csv", "人工动作及人工目标"],
            ["pred_action / pred_poi_number", "action_predictions.jsonl", "GPT 或 Qwen 的结构化动作及目标"],
            ["action_correct", "两个模型 JSONL", "动作类型是否正确"],
            ["target_correct", "两个模型 JSONL", "严格目标是否正确"],
            ["relaxed_target_correct", "两个模型 JSONL", "宽松目标是否正确"],
            ["latency_s", "Qwen JSONL", "单帧推理时延，秒"],
            ["reason / model_content", "两个模型 JSONL", "原始理由或模型结构化响应；完整保留在原始日志中"],
        ],
        font_size=8.6,
    )

    document.add_heading("附录 B  关键文件与校验值", level=1)
    source_rows = [
        source_row("paper/poi_vlm_ei_paper.md", "论文正文"),
        source_row("paper/reference_verification_report.md", "参考文献核验报告"),
        source_row("vlm_dataset_night_150/annotations.csv", "150 帧人工标签"),
        source_row("vlm_dataset_night_150/action_predictions.jsonl", "GPT-5.5 逐帧结果"),
        source_row("vlm_dataset_night_150/qwen_h800_eval/action_predictions.jsonl", "Qwen-H800 逐帧结果"),
        source_row("vlm_dataset_night_150/qwen_h800_eval/qwen_vs_gpt55_comparison.json", "模型配对对比"),
        source_row("vlm_dataset_night_150/samples.jsonl", "150 帧样本与路线元数据"),
        source_row("vlm_dataset_night/samples.jsonl", "50 帧样本元数据"),
    ]
    add_table(document, ["相对路径", "作用", "字节数", "SHA-256 前 16 位"], source_rows, font_size=7.9)
    add_paragraph(document, "SHA-256 是针对本报告生成时工作区文件内容计算的校验值；文件若被重新生成或编辑，校验值会变化。", style="Archive Note")

    # Landscape appendix for the complete per-frame ledger.
    landscape = document.add_section(WD_SECTION.NEW_PAGE)
    set_landscape(landscape)
    add_header_footer(landscape)
    document.add_heading("附录 C  150 帧逐帧测试台账", level=1)
    add_paragraph(
        document,
        "以下台账由 annotations.csv、GPT-5.5 action_predictions.jsonl 和 Qwen-H800 action_predictions.jsonl 按 sample_id 合并生成。目标列保留 POI 编号或 rotate 方向；原始 reason/model_content 仍保存在 JSONL，不在表中展开。Y 表示该字段命中，N 表示未命中。",
        style="Archive Note",
    )
    ledger_headers = ["样本", "人工动作", "人工目标", "GPT 动作", "GPT 目标", "GPT 动作命中", "GPT 目标命中", "Qwen 动作", "Qwen 目标", "Qwen 动作命中", "Qwen 目标命中", "Qwen 时延(s)"]
    ledger_rows = []
    for row in gpt_rows:
        qwen = qwen_by_id.get(row["sample_id"], {})
        ledger_rows.append(
            [
                row["sample_id"],
                row.get("label_action", "-"),
                label_target_display(row),
                row.get("pred_action", "-"),
                target_display(row),
                "Y" if row.get("action_correct") else "N",
                "Y" if row.get("target_correct") else "N",
                qwen.get("pred_action", "-"),
                target_display(qwen),
                "Y" if qwen.get("action_correct") else "N",
                "Y" if qwen.get("target_correct") else "N",
                fmt_num(float(qwen["latency_s"]), 3) if qwen.get("latency_s") is not None else "-",
            ]
        )
    add_table(document, ledger_headers, ledger_rows, font_size=6.4)

    document.add_heading("附录 D  论文与工作区版本备注", level=1)
    add_paragraph(document, "1. 论文正文中的表格、图号和方法描述按 paper/poi_vlm_ei_paper.md 版本归档。")
    add_paragraph(document, "2. 当前报告不覆盖写入论文或提交 Word，仅新增一个中文测试存档报告文件。")
    add_paragraph(document, "3. 论文中“夜间”与工作区目录的历史命名可能不完全对应；报告不擅自改写实验标签。")
    add_paragraph(document, "4. 任何对外引用结果时，应同时引用样本量和评价口径，例如“150 帧离线测试，GPT-5.5 动作准确率 138/150”。")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
